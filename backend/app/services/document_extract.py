from __future__ import annotations

from pathlib import Path

from app.services.pdf_extract import extract_text_from_pdf

ALLOWED_EXTENSIONS = {".pdf", ".pptx", ".ppt"}

ALLOWED_MIME_PREFIXES = (
    "application/pdf",
    "application/vnd.ms-powerpoint",
    "application/vnd.openxmlformats-officedocument.presentationml",
)


def is_allowed_upload(filename: str, mime_type: str | None = None) -> bool:
    lower = filename.lower()
    if any(lower.endswith(ext) for ext in ALLOWED_EXTENSIONS):
        return True
    if mime_type:
        mt = mime_type.lower()
        if any(mt.startswith(p) for p in ALLOWED_MIME_PREFIXES):
            return True
        if "presentationml" in mt or "powerpoint" in mt:
            return True
    return False


def extract_text_from_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is required: pip install python-pptx") from exc

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            text = text.strip()
            if text:
                lines.append(text)
        if lines:
            parts.append(f"## 슬라이드 {i}\n" + "\n".join(lines))
    return "\n\n".join(parts).strip()


def extract_text_from_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix == ".pptx":
        return extract_text_from_pptx(path)
    if suffix == ".ppt":
        raise ValueError(
            "구형 .ppt 형식은 지원하지 않습니다. PowerPoint에서 .pptx로 저장 후 다시 업로드해 주세요."
        )
    raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix}")
